############################################
# billing/services_sharing_exports.py
# Professionelle Export-Engine für Energy Sharing
# (PDF, Excel .xlsx, CSV, XML / ERP-Export)
############################################

import csv
import xml.etree.ElementTree as ET
from xml.dom import minidom
from io import BytesIO, StringIO
from decimal import Decimal
from datetime import datetime, date

from django.utils import timezone
from django.http import HttpResponse, Http404

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    HRFlowable,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from billing.models import CommunityMonthlyStatement, CommunityTariff, BankAccount
from accounts.models import UserProfile


# ==============================================================================
# 📄 1. PDF MONATSABRECHNUNGSNACHWEIS (Eichrechts- & Fiskal-konform)
# ==============================================================================

def generate_statement_pdf(statement: CommunityMonthlyStatement) -> HttpResponse:
    """
    Erstellt ein revisionssicheres, druckfertiges PDF für den monatlichen
    Energy-Sharing Abrechnungsnachweis mit ReportLab.
    """
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "DocTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        fontName="Helvetica-Bold",
    )
    subtitle_style = ParagraphStyle(
        "DocSubTitle",
        parent=styles["Normal"],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#475569"),
        fontName="Helvetica",
    )
    section_heading = ParagraphStyle(
        "SectionHeading",
        parent=styles["Heading2"],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#0f172a"),
        fontName="Helvetica-Bold",
        spaceBefore=8,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#334155"),
        fontName="Helvetica",
    )
    body_bold = ParagraphStyle(
        "BodyBold",
        parent=body_style,
        fontName="Helvetica-Bold",
    )
    alert_style = ParagraphStyle(
        "Alert",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#0369a1"),
        fontName="Helvetica",
    )

    story = []

    # 1. Header: Logo / System & Dokumententyp
    header_data = [
        [
            Paragraph("<b>SHAREGY</b> | Energy Sharing Clearing", title_style),
            Paragraph(f"<b>Abrechnungsnachweis</b><br/><font color='#64748b'>Nr. {statement.statement_number}</font>", subtitle_style),
        ]
    ]
    t_header = Table(header_data, colWidths=[320, 200])
    t_header.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ALIGN", (1, 0), (1, 0), "RIGHT"),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#6366f1"), spaceAfter=14))

    # 2. Metadaten (Gemeinschaft, Zeitraum, Teilnehmer, Datum)
    profile = UserProfile.objects.filter(user=statement.user).first()
    recipient_name = statement.user.get_full_name() or statement.user.username or statement.user.email
    address_str = f"{profile.street} {profile.house_number}, {profile.postal_code} {profile.city}" if profile and profile.street else "Registriertes Mitglied"

    period_str = f"{statement.period_start.strftime('%d.%m.%Y')} bis {statement.period_end.strftime('%d.%m.%Y')}"
    created_date_str = statement.created_at.strftime("%d.%m.%Y")

    meta_data = [
        [
            Paragraph("<b>Teilnehmer / Anschrift:</b>", body_bold),
            Paragraph("<b>Abrechnungsdetails:</b>", body_bold),
        ],
        [
            Paragraph(f"<b>{recipient_name}</b><br/>{address_str}<br/>E-Mail: {statement.user.email}", body_style),
            Paragraph(
                f"<b>Gemeinschaft:</b> {statement.tenant.name}<br/>"
                f"<b>Abrechnungszeitraum:</b> {period_str}<br/>"
                f"<b>Ausstellungsdatum:</b> {created_date_str}<br/>"
                f"<b>Status:</b> {statement.status.upper()}",
                body_style,
            ),
        ],
    ]
    t_meta = Table(meta_data, colWidths=[260, 260])
    t_meta.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ("PADDING", (0, 0), (-1, -1), 8),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 14))

    # 3. Mengenbilanz (OBIS 1.8.0 / 2.8.0 & Sharing)
    story.append(Paragraph("1. Mengenbilanz im 15-Minuten-Takt (kWh)", section_heading))
    mengen_data = [
        ["Messkategorie", "OBIS-Code", "Menge (kWh)", "Erläuterung"],
        ["Erzeugung (Solar)", "2.8.0", f"{statement.produced_total_kwh:,.2f} kWh", "Gesamt eingespeiste Solarenergie"],
        ["Gesamtverbrauch", "1.8.0", f"{statement.consumed_total_kwh:,.2f} kWh", "Gesamter Strombedarf der Entnahmestelle"],
        ["Geteilter Bezug (Sharing)", "1.8.0 int.", f"{statement.shared_imported_kwh:,.2f} kWh", "Aus der Gemeinschaft zeitgleich bezogen"],
        ["Geteilte Einspeisung (Sharing)", "2.8.0 int.", f"{statement.shared_exported_kwh:,.2f} kWh", "An Nachbarn in der Community geliefert"],
        ["Reststrom Netzbezug", "1.8.0 ext.", f"{statement.grid_residual_import_kwh:,.2f} kWh", "Über externen Restversorger bezogen"],
    ]
    t_mengen = Table(mengen_data, colWidths=[140, 75, 105, 200])
    t_mengen.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ALIGN", (2, 1), (2, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    story.append(t_mengen)
    story.append(Spacer(1, 14))

    # 4. Finanzielle Verrechnung
    story.append(Paragraph("2. Finanzielle Verrechnung & Vergütung (€)", section_heading))
    tariff_name = statement.tariff.name if statement.tariff else "Standard Sharing Tarif"
    sharing_price = statement.tariff.sharing_price_ct_kwh if statement.tariff else Decimal("12.00")
    payout_price = statement.tariff.producer_payout_ct_kwh if statement.tariff else Decimal("10.00")
    fee_price = statement.tariff.community_fee_ct_kwh if statement.tariff else Decimal("2.00")

    finanz_data = [
        ["Abrechnungsposition", "Menge", "Satz (Ct/kWh)", "Betrag (€)"],
        [
            f"Geteilter Strombezug ({tariff_name})",
            f"{statement.shared_imported_kwh:,.2f} kWh",
            f"{sharing_price:.2f} Ct",
            f"{statement.charge_shared_import_eur:,.2f} €",
        ],
        [
            "Einspeisevergütung Community (Gutschrift)",
            f"{statement.shared_exported_kwh:,.2f} kWh",
            f"{payout_price:.2f} Ct",
            f"-{statement.credit_shared_export_eur:,.2f} €",
        ],
        [
            "Community-Betriebsumlage (Software & Clearing)",
            f"{statement.shared_imported_kwh:,.2f} kWh",
            f"{fee_price:.2f} Ct",
            f"{statement.community_fee_eur:,.2f} €",
        ],
    ]
    t_finanz = Table(finanz_data, colWidths=[220, 100, 100, 100])
    t_finanz.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#475569")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
    ]))
    story.append(t_finanz)
    story.append(Spacer(1, 12))

    # 5. Saldo Highlight Box
    net_val = float(statement.net_balance_eur)
    is_payout = net_val >= 0
    box_color = colors.HexColor("#10b981") if is_payout else colors.HexColor("#ef4444")
    saldo_label = "GUTSCHRIFT / AUSZAHLUNGSBETRAG:" if is_payout else "ZAHLBETRAG / FORDERUNG:"
    saldo_text = f"+{statement.net_balance_eur:,.2f} €" if is_payout else f"{statement.net_balance_eur:,.2f} €"

    saldo_data = [
        [
            Paragraph(f"<b>{saldo_label}</b>", ParagraphStyle("SLabel", parent=body_bold, fontSize=11, textColor=box_color)),
            Paragraph(f"<b>{saldo_text}</b>", ParagraphStyle("SVal", parent=body_bold, fontSize=16, textColor=box_color, alignment=2)),
        ]
    ]
    t_saldo = Table(saldo_data, colWidths=[320, 200])
    t_saldo.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#ecfdf5") if is_payout else colors.HexColor("#fef2f2")),
        ("PADDING", (0, 0), (-1, -1), 10),
        ("BOX", (0, 0), (-1, -1), 1.5, box_color),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(t_saldo)
    story.append(Spacer(1, 14))

    # 6. Rechtliche Hinweise & Fußzeile
    legal_text = (
        "<b>Hinweise zur Abrechnung:</b><br/>"
        "• Dieser Abrechnungsnachweis dient der internen Verrechnung innerhalb der Energiegemeinschaft gem. § 42b EnWG.<br/>"
        "• Der verbleibende Reststrombedarf wird separat von Ihrem gewählten Reststromlieferanten abgerechnet.<br/>"
        "• Bei Guthaben erfolgt die Überweisung auf Ihr hinterlegtes Bankkonto. Bei Zahlbeträgen erfolgt der Einzug zum Monatsende.<br/>"
        "• Eichrechtlich validiert über 15-Minuten-Messwerte (OBIS 1.8.0 / 2.8.0) der zertifizierten Smart-Meter-Gateways."
    )
    story.append(Paragraph(legal_text, body_style))
    story.append(Spacer(1, 16))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cbd5e1"), spaceAfter=8))
    story.append(Paragraph(f"Sharegy Energy Sharing Engine | Tenant: {statement.tenant.name} | Dokument-ID: {statement.id}", subtitle_style))

    doc.build(story)
    pdf_value = buffer.getvalue()
    buffer.close()

    filename = f"Abrechnungsnachweis_{statement.statement_number}.pdf"
    response = HttpResponse(pdf_value, content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# 📊 2. CSV EXPORT FÜR HAUSVERWALTUNGEN & ERP
# ==============================================================================

def export_statements_csv(statements_qs, tenant_name: str = "Community") -> HttpResponse:
    """
    Exportiert Monatsabrechnungen als formatiertes CSV (UTF-8 mit BOM, Semikolon-getrennt)
    zur direkten Weiterverarbeitung in Excel oder ERP-Systemen.
    """
    string_io = StringIO()
    # UTF-8 BOM für fehlerfreie deutsche Umlaute in Excel
    string_io.write("\ufeff")
    writer = csv.writer(string_io, delimiter=";", lineterminator="\r\n")

    # Header-Metadaten
    writer.writerow(["# Sharegy Energy Sharing - Abrechnungsdaten-Export"])
    writer.writerow(["# Gemeinschaft:", tenant_name])
    writer.writerow(["# Exportdatum:", timezone.now().strftime("%d.%m.%Y %H:%M")])
    writer.writerow(["# Anzahl Datensaetze:", str(statements_qs.count())])
    writer.writerow([])

    # Spaltenüberschriften
    writer.writerow([
        "Abrechnungsnummer",
        "Abrechnungsmonat",
        "Zeitraum_Start",
        "Zeitraum_Ende",
        "Teilnehmer_Email",
        "Erzeugung_kWh",
        "Verbrauch_kWh",
        "Geteilt_Import_kWh",
        "Geteilt_Export_kWh",
        "Reststrom_Netzbezug_kWh",
        "Kosten_Geteilter_Bezug_EUR",
        "Verguetung_Einspeisung_EUR",
        "Community_Umlage_EUR",
        "Netto_Saldo_EUR",
        "Art_Saldo",
        "Status",
        "Finalisiert_Am",
    ])

    for s in statements_qs:
        net_val = float(s.net_balance_eur)
        art_saldo = "Gutschrift" if net_val >= 0 else "Forderung"
        writer.writerow([
            s.statement_number,
            s.period_start.strftime("%Y-%m"),
            s.period_start.strftime("%d.%m.%Y"),
            s.period_end.strftime("%d.%m.%Y"),
            s.user.email,
            str(s.produced_total_kwh).replace(".", ","),
            str(s.consumed_total_kwh).replace(".", ","),
            str(s.shared_imported_kwh).replace(".", ","),
            str(s.shared_exported_kwh).replace(".", ","),
            str(s.grid_residual_import_kwh).replace(".", ","),
            str(s.charge_shared_import_eur).replace(".", ","),
            str(s.credit_shared_export_eur).replace(".", ","),
            str(s.community_fee_eur).replace(".", ","),
            str(s.net_balance_eur).replace(".", ","),
            art_saldo,
            s.status,
            s.finalized_at.strftime("%d.%m.%Y %H:%M") if s.finalized_at else "",
        ])

    csv_data = string_io.getvalue().encode("utf-8")
    filename = f"Sharegy_Abrechnungsdaten_{tenant_name.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.csv"
    response = HttpResponse(csv_data, content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# 📗 3. EXCEL .XLSX EXPORT MIT FORMELN & STYLING
# ==============================================================================

def export_statements_xlsx(statements_qs, tenant_name: str = "Community") -> HttpResponse:
    """
    Erstellt eine formatierte Excel-Arbeitsmappe (.xlsx) mit Summenformeln
    und Farbcodierung für Buchhaltung & Hausverwaltung.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Abrechnungsnachweise"

    # Styling-Definitionen
    font_header = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=9)
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_green = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
    fill_red = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Titelzeilen
    ws.append([f"Sharegy Energy Sharing - Monatsabrechnungen: {tenant_name}"])
    ws.append([f"Exportiert am: {timezone.now().strftime('%d.%m.%Y %H:%M')} Uhr"])
    ws.append([])
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="0F172A")
    ws["A2"].font = Font(name="Arial", size=9, color="64748B")

    headers = [
        "Abrechnungs-Nr.",
        "Monat",
        "Teilnehmer (E-Mail)",
        "Erzeugt (kWh)",
        "Verbraucht (kWh)",
        "Geteilt Import (kWh)",
        "Geteilt Export (kWh)",
        "Reststrom (kWh)",
        "Kosten Bezug (€)",
        "Gutschrift Erzeugung (€)",
        "Community-Umlage (€)",
        "Netto-Saldo (€)",
        "Status",
    ]
    ws.append(headers)
    header_row_idx = 4

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_start = 5
    current_row = row_start
    for s in statements_qs:
        ws.append([
            s.statement_number,
            s.period_start.strftime("%Y-%m"),
            s.user.email,
            float(s.produced_total_kwh),
            float(s.consumed_total_kwh),
            float(s.shared_imported_kwh),
            float(s.shared_exported_kwh),
            float(s.grid_residual_import_kwh),
            float(s.charge_shared_import_eur),
            float(s.credit_shared_export_eur),
            float(s.community_fee_eur),
            float(s.net_balance_eur),
            s.status,
        ])

        # Styling der Datenzeile
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=current_row, column=col_idx)
            c.font = font_regular
            c.border = thin_border
            if col_idx in [4, 5, 6, 7, 8]:
                c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
            elif col_idx in [9, 10, 11, 12]:
                c.number_format = "#,##0.00 €"
                c.alignment = Alignment(horizontal="right")
                if col_idx == 12:  # Netto-Saldo
                    if float(s.net_balance_eur) >= 0:
                        c.fill = fill_green
                    else:
                        c.fill = fill_red

        current_row += 1

    # Summenzeile falls Daten vorhanden
    if current_row > row_start:
        summary_row = current_row
        ws.cell(row=summary_row, column=1, value="GESAMT").font = font_bold
        ws.cell(row=summary_row, column=1).border = thin_border

        for col_idx, col_letter in [
            (4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H"),
            (9, "I"), (10, "J"), (11, "K"), (12, "L")
        ]:
            cell = ws.cell(row=summary_row, column=col_idx)
            cell.value = f"=SUM({col_letter}{row_start}:{col_letter}{summary_row - 1})"
            cell.font = font_bold
            cell.border = thin_border
            cell.number_format = "#,##0.00 €" if col_idx >= 9 else "#,##0.00"

    # Spaltenbreiten optimieren
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = BytesIO()
    wb.save(buffer)
    excel_data = buffer.getvalue()
    buffer.close()

    filename = f"Sharegy_Monatsabrechnungen_{tenant_name.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(excel_data, content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# 📦 4. XML / ERP EXPORT (Standardisiertes Buchhaltungsformat)
# ==============================================================================

def export_statements_xml(statements_qs, tenant_name: str = "Community") -> HttpResponse:
    """
    Generiert standardisiertes XML für ERP- und Buchhaltungssysteme.
    """
    root = ET.Element("EnergySharingSettlementExport")
    root.set("version", "1.0")
    root.set("system", "Sharegy Clearing Engine")
    root.set("exportedAt", timezone.now().isoformat())

    tenant_el = ET.SubElement(root, "Community")
    tenant_el.set("name", tenant_name)

    statements_el = ET.SubElement(root, "Statements")
    statements_el.set("count", str(statements_qs.count()))

    for s in statements_qs:
        stmt_el = ET.SubElement(statements_el, "Statement")
        stmt_el.set("number", s.statement_number)
        stmt_el.set("periodStart", s.period_start.isoformat())
        stmt_el.set("periodEnd", s.period_end.isoformat())
        stmt_el.set("status", s.status)

        member_el = ET.SubElement(stmt_el, "Member")
        member_el.set("userId", str(s.user.id))
        member_el.set("email", s.user.email)

        energy_el = ET.SubElement(stmt_el, "EnergyQuantities", unit="kWh")
        ET.SubElement(energy_el, "ProducedTotal").text = f"{s.produced_total_kwh:.3f}"
        ET.SubElement(energy_el, "ConsumedTotal").text = f"{s.consumed_total_kwh:.3f}"
        ET.SubElement(energy_el, "SharedImport").text = f"{s.shared_imported_kwh:.3f}"
        ET.SubElement(energy_el, "SharedExport").text = f"{s.shared_exported_kwh:.3f}"
        ET.SubElement(energy_el, "GridResidualImport").text = f"{s.grid_residual_import_kwh:.3f}"

        finance_el = ET.SubElement(stmt_el, "FinancialSettlement", currency="EUR")
        ET.SubElement(finance_el, "ChargeSharedImport").text = f"{s.charge_shared_import_eur:.2f}"
        ET.SubElement(finance_el, "CreditSharedExport").text = f"{s.credit_shared_export_eur:.2f}"
        ET.SubElement(finance_el, "CommunityFee").text = f"{s.community_fee_eur:.2f}"
        ET.SubElement(finance_el, "NetBalance").text = f"{s.net_balance_eur:.2f}"

    xml_str = minidom.parseString(ET.tostring(root, encoding="utf-8")).toprettyxml(indent="  ", encoding="utf-8")
    filename = f"Sharegy_Abrechnungsdaten_{tenant_name.replace(' ', '_')}_{timezone.now().strftime('%Y%m%d')}.xml"
    response = HttpResponse(xml_str, content_type="application/xml; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
