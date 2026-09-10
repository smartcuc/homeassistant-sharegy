#####################################
# energy/services/export_manager.py
#####################################

import csv
import json
from io import BytesIO, StringIO
from zoneinfo import ZoneInfo
from django.utils import timezone
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from energy.services.balance import get_energy_balance


def export_energy_balance(user, period="today", start_date=None, end_date=None, export_format="xlsx") -> HttpResponse:
    data = get_energy_balance(user, period=period, start_date=start_date, end_date=end_date)
    period_label = data.get("period_label", period)
    kpis = data.get("kpis", {})
    submeters = data.get("submeters", [])
    timeseries = data.get("charts", {}).get("timeseries", [])
    insights = data.get("insights", [])

    home = user.homes.first() if hasattr(user, "homes") else None
    tz_name = home.timezone if home and home.timezone else "Europe/Berlin"
    export_ts_str = timezone.now().astimezone(ZoneInfo(tz_name)).strftime("%d.%m.%Y %H:%M")

    safe_period = str(period_label).replace(".", "-").replace("/", "-").replace(" ", "_").replace(":", "-")

    # 1. JSON Export
    if export_format == "json":
        payload = {
            "meta": {
                "system": "Sharegy EMS Cloud",
                "user": user.username,
                "exported_at": export_ts_str,
                "period": period,
                "period_label": period_label,
            },
            "kpis": kpis,
            "submeters": submeters,
            "timeseries": timeseries,
            "insights": insights,
        }
        resp = HttpResponse(json.dumps(payload, indent=2, ensure_ascii=False), content_type="application/json; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="sharegy_energiebilanz_{safe_period}.json"'
        return resp

    # 2. CSV Export
    elif export_format == "csv":
        string_io = StringIO()
        writer = csv.writer(string_io, delimiter=";", lineterminator="\r\n")

        writer.writerow(["# Sharegy EMS - Energiebilanz & Verbrauchsbericht"])
        writer.writerow(["# Zeitraum:", period_label])
        writer.writerow(["# Exportdatum:", export_ts_str])
        writer.writerow(["# Benutzer:", user.username])
        writer.writerow([])

        writer.writerow(["--- KPI ZUSAMMENFASSUNG ---"])
        writer.writerow(["Kennzahl", "Wert", "Einheit"])
        writer.writerow(["Solar-Erzeugung (PV)", str(kpis.get("pv_generation_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Gesamter Hausverbrauch", str(kpis.get("house_consumption_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Netzbezug", str(kpis.get("grid_import_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Netzeinspeisung", str(kpis.get("grid_export_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Batterieladung", str(kpis.get("battery_charge_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Batterieentladung", str(kpis.get("battery_discharge_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Solarer Eigenverbrauch", str(kpis.get("direct_consumption_kwh", 0)).replace(".", ","), "kWh"])
        writer.writerow(["Autarkiegrad", str(kpis.get("autarky_rate", 0)).replace(".", ","), "%"])
        writer.writerow(["Eigenverbrauchsquote", str(kpis.get("self_consumption_rate", 0)).replace(".", ","), "%"])
        writer.writerow(["Stromkosten-Einsparung", str(kpis.get("savings_eur", 0)).replace(".", ","), "EUR"])
        writer.writerow(["Einspeisevergütung", str(kpis.get("feed_in_revenue_eur", 0)).replace(".", ","), "EUR"])
        writer.writerow(["Netzstromkosten", str(kpis.get("grid_costs_eur", 0)).replace(".", ","), "EUR"])
        writer.writerow(["Finanzieller Netto-Vorteil", str(kpis.get("net_benefit_eur", 0)).replace(".", ","), "EUR"])
        writer.writerow(["CO2-Einsparung", str(kpis.get("co2_saved_kg", 0)).replace(".", ","), "kg"])
        writer.writerow([])

        if submeters:
            writer.writerow(["--- SUB-METERING & VERBRAUCHER ---"])
            writer.writerow(["Kategorie / Gerät", "Verbrauch (kWh)", "Anteil (%)", "Solaranteil (%)", "Kosten (EUR)", "Ersparnis (EUR)"])
            for sm in submeters:
                writer.writerow([
                    sm.get("name", ""),
                    str(sm.get("consumption_kwh", 0)).replace(".", ","),
                    str(sm.get("share_pct", 0)).replace(".", ","),
                    str(sm.get("solar_share_pct", 0)).replace(".", ","),
                    str(sm.get("cost_eur", 0)).replace(".", ","),
                    str(sm.get("savings_eur", 0)).replace(".", ","),
                ])
            writer.writerow([])

        if timeseries:
            writer.writerow(["--- ZEITREIHEN-DETAILS ---"])
            writer.writerow(["Zeitpunkt", "Solar PV (kWh)", "Hausverbrauch (kWh)", "Batterie Entladung (kWh)", "Netzbezug (kWh)", "Netzeinspeisung (kWh)"])
            for row in timeseries:
                writer.writerow([
                    row.get("time", ""),
                    str(row.get("pv", 0)).replace(".", ","),
                    str(row.get("load", 0)).replace(".", ","),
                    str(row.get("battery_discharge", 0)).replace(".", ","),
                    str(row.get("grid_import", 0)).replace(".", ","),
                    str(row.get("grid_export", 0)).replace(".", ","),
                ])

        csv_content = b"\xef\xbb\xbf" + string_io.getvalue().encode("utf-8")
        resp = HttpResponse(csv_content, content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="sharegy_energiebilanz_{safe_period}.csv"'
        return resp

    # 3. Excel Export (.xlsx)
    elif export_format == "xlsx":
        wb = Workbook()
        ws_kpi = wb.active
        ws_kpi.title = "Energiebilanz Übersicht"
        ws_kpi.views.sheetView[0].showGridLines = True

        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        sub_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
        title_font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        ws_kpi["A1"] = "Sharegy EMS - Energie- & Kostenbilanz"
        ws_kpi["A1"].font = title_font
        ws_kpi["A2"] = f"Zeitraum: {period_label} | Exportiert am: {export_ts_str} | Benutzer: {user.username}"
        ws_kpi["A2"].font = Font(name="Calibri", size=10, italic=True, color="64748B")

        headers = ["Kennzahl", "Menge / Wert", "Einheit", "Finanzieller Effekt (€)"]
        for col_idx, h in enumerate(headers, 1):
            cell = ws_kpi.cell(row=4, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="left", vertical="center")

        kpi_rows = [
            ("Solar-Erzeugung (PV)", kpis.get("pv_generation_kwh", 0), "kWh", f"+ {kpis.get('savings_eur', 0):.2f} € (Eigenverbrauch)"),
            ("Gesamter Hausverbrauch", kpis.get("house_consumption_kwh", 0), "kWh", f"- {kpis.get('grid_costs_eur', 0):.2f} € (Netzkosten)"),
            ("Solarer Direktverbrauch", kpis.get("direct_consumption_kwh", 0), "kWh", "-"),
            ("Batteriespeicher Ladung", kpis.get("battery_charge_kwh", 0), "kWh", "-"),
            ("Batteriespeicher Entladung", kpis.get("battery_discharge_kwh", 0), "kWh", "-"),
            ("Netzbezug (Import)", kpis.get("grid_import_kwh", 0), "kWh", f"- {kpis.get('grid_costs_eur', 0):.2f} €"),
            ("Netzeinspeisung (Export)", kpis.get("grid_export_kwh", 0), "kWh", f"+ {kpis.get('feed_in_revenue_eur', 0):.2f} €"),
            ("Autarkiegrad", kpis.get("autarky_rate", 0), "%", "Solare Unabhängigkeit"),
            ("Eigenverbrauchsquote", kpis.get("self_consumption_rate", 0), "%", "PV-Nutzungsgrad"),
            ("CO₂-Einsparung", kpis.get("co2_saved_kg", 0), "kg", f"≈ {kpis.get('trees_equivalent', 0)} Bäume"),
            ("FINANZIELLER NETTO-VORTEIL", kpis.get("net_benefit_eur", 0), "EUR", "Ersparnis + Vergütung - Netzkosten"),
        ]

        curr_row = 5
        for label, val, unit, fin in kpi_rows:
            ws_kpi.cell(row=curr_row, column=1, value=label).border = thin_border
            ws_kpi.cell(row=curr_row, column=2, value=val).border = thin_border
            ws_kpi.cell(row=curr_row, column=3, value=unit).border = thin_border
            ws_kpi.cell(row=curr_row, column=4, value=fin).border = thin_border
            if label.startswith("FINANZIELLER"):
                for col in range(1, 5):
                    ws_kpi.cell(row=curr_row, column=col).font = bold_font
                    ws_kpi.cell(row=curr_row, column=col).fill = sub_fill
            curr_row += 1

        if submeters:
            curr_row += 2
            ws_kpi.cell(row=curr_row, column=1, value="Aufteilung nach Verbrauchern (Sub-Metering)").font = Font(size=13, bold=True, color="1E1B4B")
            curr_row += 1
            sm_headers = ["Verbraucher", "Verbrauch (kWh)", "Anteil (%)", "Solaranteil (%)", "Kosten (€)", "Ersparnis (€)"]
            for col_idx, h in enumerate(sm_headers, 1):
                cell = ws_kpi.cell(row=curr_row, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
            curr_row += 1
            for sm in submeters:
                ws_kpi.cell(row=curr_row, column=1, value=sm.get("name")).border = thin_border
                ws_kpi.cell(row=curr_row, column=2, value=sm.get("consumption_kwh")).border = thin_border
                ws_kpi.cell(row=curr_row, column=3, value=f"{sm.get('share_pct') or 0}%").border = thin_border
                ws_kpi.cell(row=curr_row, column=4, value=f"{sm.get('solar_share_pct') or 0}%").border = thin_border
                ws_kpi.cell(row=curr_row, column=5, value=f"{sm.get('cost_eur') or 0} €").border = thin_border
                ws_kpi.cell(row=curr_row, column=6, value=f"{sm.get('savings_eur') or 0} €").border = thin_border
                curr_row += 1

        for col in ws_kpi.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_kpi.column_dimensions[col_letter].width = max(max_len + 3, 14)

        if timeseries:
            ws_ts = wb.create_sheet(title="Zeitreihen-Details")
            ws_ts.views.sheetView[0].showGridLines = True
            ts_headers = ["Zeitpunkt", "Solar PV (kWh)", "Hauslast (kWh)", "Batterie Entladung (kWh)", "Netzbezug (kWh)", "Netzeinspeisung (kWh)"]
            for col_idx, h in enumerate(ts_headers, 1):
                cell = ws_ts.cell(row=1, column=col_idx, value=h)
                cell.font = header_font
                cell.fill = header_fill
            ts_row = 2
            for row in timeseries:
                ws_ts.cell(row=ts_row, column=1, value=row.get("time")).border = thin_border
                ws_ts.cell(row=ts_row, column=2, value=row.get("pv", 0)).border = thin_border
                ws_ts.cell(row=ts_row, column=3, value=row.get("load", 0)).border = thin_border
                ws_ts.cell(row=ts_row, column=4, value=row.get("battery_discharge", 0)).border = thin_border
                ws_ts.cell(row=ts_row, column=5, value=row.get("grid_import", 0)).border = thin_border
                ws_ts.cell(row=ts_row, column=6, value=row.get("grid_export", 0)).border = thin_border
                ts_row += 1

            for col in ws_ts.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws_ts.column_dimensions[col_letter].width = max(max_len + 3, 16)

        resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = f'attachment; filename="sharegy_energiebilanz_{safe_period}.xlsx"'
        wb.save(resp)
        return resp

    # 4. PDF Export
    else:
        buffer = BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36,
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "SharegyTitle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#1E1B4B"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "SharegySubtitle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=14,
        )
        section_style = ParagraphStyle(
            "SharegySection",
            parent=styles["Heading2"],
            fontSize=12,
            leading=15,
            textColor=colors.HexColor("#4F46E5"),
            spaceBefore=12,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "SharegyBody",
            parent=styles["Normal"],
            fontSize=9,
            leading=12,
            textColor=colors.HexColor("#334155"),
        )

        elements = []

        elements.append(Paragraph("☀️ Sharegy EMS — Energie- & Kostenbericht", title_style))
        elements.append(Paragraph(f"Zeitraum: <b>{period_label}</b> &nbsp;|&nbsp; Erstellt am: {export_ts_str} &nbsp;|&nbsp; Kunde: {user.username}", subtitle_style))
        elements.append(Spacer(1, 6))

        elements.append(Paragraph("1. Kennzahlen & Energiefluss", section_style))
        kpi_table_data = [
            [Paragraph("<b>Kennzahl</b>", body_style), Paragraph("<b>Menge</b>", body_style), Paragraph("<b>Finanzielle Bewertung / Nutzen</b>", body_style)],
            ["Solar-Erzeugung (PV)", f"{kpis.get('pv_generation_kwh', 0):.1f} kWh", f"+ {kpis.get('savings_eur', 0):.2f} € Stromkosteneinsparung"],
            ["Gesamter Hausverbrauch", f"{kpis.get('house_consumption_kwh', 0):.1f} kWh", f"Deckungsgrad: {kpis.get('autarky_rate', 0)} % autark"],
            ["Solarer Eigenverbrauch", f"{kpis.get('direct_consumption_kwh', 0):.1f} kWh", f"Eigenverbrauchsquote: {kpis.get('self_consumption_rate', 0)} %"],
            ["Batteriespeicher (Ladung/Entladung)", f"{kpis.get('battery_charge_kwh', 0):.1f} / {kpis.get('battery_discharge_kwh', 0):.1f} kWh", "Pufferung für Abend- & Nachtstunden"],
            ["Netzbezug (Zukauf)", f"{kpis.get('grid_import_kwh', 0):.1f} kWh", f"- {kpis.get('grid_costs_eur', 0):.2f} € Strombezugskosten"],
            ["Netzeinspeisung (Überschuss)", f"{kpis.get('grid_export_kwh', 0):.1f} kWh", f"+ {kpis.get('feed_in_revenue_eur', 0):.2f} € Einspeisevergütung"],
            ["CO₂-Vermeidung", f"{kpis.get('co2_saved_kg', 0):.1f} kg CO₂", f"Ökologischer Beitrag (≈ {kpis.get('trees_equivalent', 0)} Bäume)"],
            ["FINANZIELLER NETTO-VORTEIL", f"{kpis.get('net_benefit_eur', 0):.2f} €", "Ersparnis + Vergütung - Netzkosten"],
        ]

        t = Table(kpi_table_data, colWidths=[170, 120, 230])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.HexColor("#F8FAFC"), colors.white]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EEF2FF")),
            ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 10))

        if submeters:
            elements.append(Paragraph("2. Verbraucher-Aufschlüsselung (Sub-Metering)", section_style))
            sm_data = [[
                Paragraph("<b>Verbraucher</b>", body_style),
                Paragraph("<b>Verbrauch</b>", body_style),
                Paragraph("<b>Anteil</b>", body_style),
                Paragraph("<b>Solaranteil</b>", body_style),
                Paragraph("<b>Kosten (€)</b>", body_style),
            ]]
            for sm in submeters:
                sm_data.append([
                    sm.get("name", ""),
                    f"{sm.get('consumption_kwh', 0):.1f} kWh",
                    f"{sm.get('share_pct', 0):.1f} %",
                    f"{sm.get('solar_share_pct', 0):.1f} %",
                    f"{sm.get('cost_eur', 0):.2f} €",
                ])
            sm_t = Table(sm_data, colWidths=[180, 85, 85, 85, 85])
            sm_t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#6366F1")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
                ("LINEBELOW", (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ]))
            elements.append(sm_t)
            elements.append(Spacer(1, 10))

        if insights:
            elements.append(Paragraph("3. Zusammenfassung & Hinweise", section_style))
            for ins in insights:
                elements.append(Paragraph(f"• {ins}", body_style))
                elements.append(Spacer(1, 2))

        doc.build(elements)
        resp = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        resp["Content-Disposition"] = f'attachment; filename="sharegy_energiebilanz_{safe_period}.pdf"'
        return resp
